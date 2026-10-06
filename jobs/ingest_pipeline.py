"""Phase 2 — the ingestion pipeline (the credibility-critical core).

For each inbound file:
  1. parse it (title block + header row + data) — unreadable -> UNUSABLE
  2. fingerprint against the counterparty template registry -> one of the four states
  3. two-path extraction: deterministic (template field_map on the actual headers) +
     AI (dictionary-constrained ai_query proposing header->logical mapping)
  4. reconcile the two paths field-by-field via the governed fn_reconcile_status
  5. land (known & unchanged) with assertion-level provenance + typed projection, or
     quarantine (changed / unknown / unusable) with the reconciliation detail + proposed remap
Every path writes an append-only audit event. Supersession is a status flip; downstream
reads go through v_*_latest.

Importable (app / job / MCP all call the same functions). Runs against serverless.
Run standalone:  uv run --with databricks-sdk,openpyxl python jobs/ingest_pipeline.py
"""
from __future__ import annotations

import io
import json
import os
import sys
import uuid
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from lib import load_config, load_tenant, sql, insert_rows, client  # noqa: E402

HEADER_ROW = 5
LOB_CODE = {"Commercial Property": "PROP", "General Liability": "GL",
            "Commercial Motor": "MOTOR", "Marine Cargo": "MARINE",
            "Professional Indemnity": "PI"}
CAUSE_CODE = {"Escape of water": "EOW", "Fire": "FIRE", "Theft": "THEFT", "Storm": "STORM",
              "Accidental damage": "AD", "Liability": "LIAB", "Impact": "IMP"}
FAMILY_KEY = {"premium_bordereau": "premium", "claims_bordereau": "claims"}

NODE_COLS = [("node_id", "str"), ("label", "str"), ("properties", "str"),
             ("source_document_id", "str"), ("extraction_path", "str"), ("confidence", "float"),
             ("template_version", "str"), ("dictionary_version", "str"), ("anchor", "str"),
             ("tenant", "str"), ("created_at", "ts_now")]
EDGE_COLS = [("edge_id", "str"), ("source_id", "str"), ("target_id", "str"),
             ("relationship_type", "str"), ("source_document_id", "str"),
             ("extraction_path", "str"), ("confidence", "float"), ("template_version", "str"),
             ("dictionary_version", "str"), ("anchor", "str"), ("tenant", "str"),
             ("created_at", "ts_now")]


# ------------------------------------------------------------------- helpers
def _short(doc_id):
    return doc_id.split("-")[-1][:8]


def _col_letter(idx):  # 1-indexed
    from openpyxl.utils import get_column_letter
    return get_column_letter(idx)


def _month_date(m):  # "2026-07" -> "2026-07-01"
    return f"{m}-01"


def audit(cfg, event_type, entity_type, entity_id, detail="", actor="pipeline"):
    F = cfg["full_schema"]
    sql(f"""INSERT INTO {F}.audit_event VALUES
            (:id,:et,:ent,:eid,:d,:a,current_timestamp())""",
        params={"id": f"AE-{uuid.uuid4().hex[:10]}", "et": event_type, "ent": entity_type,
                "eid": entity_id, "d": detail, "a": actor}, cfg=cfg)


# ------------------------------------------------------------------- parse
class Unusable(Exception):
    pass


def parse_file(cfg, path):
    """Return {name, ref, month, family, headers, rows}. Raise Unusable if not a table."""
    fname = os.path.basename(path)
    stem = fname.rsplit(".", 1)[0]
    ext = fname.rsplit(".", 1)[-1].lower()
    parts = stem.split("_")
    ref = parts[0] if parts else ""
    family_token = parts[1] if len(parts) > 1 else ""
    month_token = parts[2] if len(parts) > 2 else ""
    family = {"premium": "premium_bordereau", "claims": "claims_bordereau"}.get(family_token, family_token)

    if ext != "xlsx":
        raise Unusable(f"not a spreadsheet (.{ext}) — no extractable table")

    from openpyxl import load_workbook
    data = client(cfg).files.download(path).contents.read()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    grid = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    if len(grid) < HEADER_ROW + 1:
        raise Unusable("too few rows — no data table")
    name = (grid[0][0] if grid[0] else None) or ref
    month = (grid[2][0] if len(grid) > 2 and grid[2] else None) or month_token
    header = [h for h in grid[HEADER_ROW - 1] if h is not None]
    if len(header) < 4:
        raise Unusable("no recognisable header row")
    rows = []
    for r in grid[HEADER_ROW:]:
        if not any(v is not None for v in r):
            continue
        rows.append({header[i]: r[i] for i in range(len(header)) if i < len(r)})
    return {"name": str(name), "ref": str(ref), "month": str(month), "family": family,
            "family_token": family_token, "headers": header, "rows": rows}


# ------------------------------------------------------------------- identify / extract
def lookup_template(cfg, counterparty, family):
    F = cfg["full_schema"]
    r = sql(f"""SELECT template_version, expected_anchors, field_map, field_count
                FROM {F}.v_template_latest
                WHERE counterparty=:c AND doc_family=:f""",
            params={"c": counterparty, "f": family}, cfg=cfg)
    if not r:
        return None
    v, anchors, fmap, n = r[0]
    return {"version": v, "expected_anchors": json.loads(anchors),
            "field_map": json.loads(fmap), "field_count": int(n)}


def det_extract(field_map, headers, rows):
    """Deterministic: for each logical field, is a template label present? map values.

    A field_map value may hold pipe-separated alternatives (e.g. 'Gross Premium|GWP')
    learned across template versions — the first alternative present in the headers wins.
    """
    hset = set(headers)
    out = {}
    for logical, spec in field_map.items():
        alts = [a.strip() for a in str(spec).split("|")]
        matched = next((a for a in alts if a in hset), None)
        out[logical] = {"label": spec, "matched": matched, "present": matched is not None}
    mapped = []
    for row in rows:
        m = {}
        for logical, info in out.items():
            if info["present"]:
                m[logical] = row.get(info["matched"])
        mapped.append(m)
    return out, mapped


def ai_extract(cfg, headers, sample_rows, logical_fields, endpoint=None):
    """AI path: dictionary-constrained ai_query proposing header->logical mapping.

    `endpoint` overrides the configured extraction model (used by the eval harness for
    model-swap comparison)."""
    prompt = (
        "You map spreadsheet columns to canonical insurance bordereau fields. "
        "Allowed canonical fields (map ONLY to these, never invent): "
        + ", ".join(logical_fields) + ".\n"
        "Given the column headers and sample rows below, return STRICT JSON: an object whose "
        "keys are the canonical field names and whose values are objects "
        '{"matched_header": <exact header string or null>, "confidence": <0..1>}. '
        "Use null when no column matches. Return JSON only, no prose.\n\n"
        f"HEADERS: {json.dumps(headers)}\n"
        f"SAMPLE_ROWS: {json.dumps(sample_rows[:2], default=str)}\n"
    )
    ep = endpoint or cfg["models"]["extract"]
    try:
        out = sql("SELECT ai_query(:ep, :p)", params={"ep": ep, "p": prompt}, cfg=cfg)
        txt = out[0][0] if out and out[0] else ""
        txt = txt.strip()
        if txt.startswith("```"):
            txt = txt.split("```")[1].lstrip("json").strip() if "```" in txt else txt
        parsed = json.loads(txt)
        return {k: {"matched_header": v.get("matched_header"),
                    "confidence": float(v.get("confidence", 0) or 0)}
                for k, v in parsed.items() if k in logical_fields}
    except Exception as e:  # AI unavailable -> reconciliation falls back to deterministic
        print(f"    (ai_query fallback: {e})")
        return {}


def reconcile(cfg, det_map, ai_map, logical_fields):
    """Field-level reconciliation via the governed fn_reconcile_status."""
    F = cfg["full_schema"]
    fields, proposed_remap = {}, []
    n_differs = 0
    for logical in logical_fields:
        det_label = det_map.get(logical, {}).get("matched") if det_map.get(logical, {}).get("present") else None
        ai_info = ai_map.get(logical) or {}
        ai_header = ai_info.get("matched_header")
        status = sql(f"SELECT {F}.fn_reconcile_status(:d,:a,0.02)",
                     params={"d": det_label, "a": ai_header}, cfg=cfg)[0][0]
        # If AI path is unavailable entirely, don't punish present deterministic fields.
        if not ai_map and det_map.get(logical, {}).get("present"):
            status = "agree"
        fields[logical] = {"det_label": det_label, "det_present": bool(det_label),
                           "ai_header": ai_header, "ai_confidence": ai_info.get("confidence"),
                           "status": status}
        if status in ("differs", "det-only"):
            n_differs += 1
        # a field the template expected but the file no longer labels that way, yet AI found it
        if not det_map.get(logical, {}).get("present") and ai_header:
            proposed_remap.append({"field": logical,
                                   "expected_label": det_map.get(logical, {}).get("label"),
                                   "found_label": ai_header})
    return {"fields": fields, "proposed_remap": proposed_remap,
            "summary": {"expected": len(logical_fields),
                        "present": sum(1 for f in det_map.values() if f["present"]),
                        "differs": n_differs}}


# ------------------------------------------------------------------- land / quarantine
def _supersede(cfg, counterparty, family, period):
    F = cfg["full_schema"]
    prior = sql(f"""SELECT doc_id FROM {F}.source_document
                    WHERE counterparty=:c AND doc_family=:f AND reporting_period=:p
                    AND status='active'""",
                params={"c": counterparty, "f": family, "p": period}, cfg=cfg)
    for (doc_id,) in prior:
        sql(f"UPDATE {F}.source_document SET status='superseded' WHERE doc_id=:d",
            params={"d": doc_id}, cfg=cfg)
        audit(cfg, "superseded", "source_document", doc_id,
              detail=f"superseded by newer {family} for {counterparty} {period}")
    return [p[0] for p in prior]


def _land_common(cfg, parsed, template, recon, doc_id):
    """Supersede prior active, insert the active source_document, and the shared
    Coverholder + ReportingPeriod nodes. Returns landing context."""
    F = cfg["full_schema"]
    name, ref, month, family = parsed["name"], parsed["ref"], parsed["month"], parsed["family"]
    superseded = _supersede(cfg, name, family, month)
    sql(f"""INSERT INTO {F}.source_document VALUES
            (:doc,'bordereaux',:c,:f,:fn,:sp,'active','known_unchanged',:tv,'v1',:rd,:rp,
             current_timestamp(),NULL,NULL)""",
        params={"doc": doc_id, "c": name, "f": family, "fn": parsed["file_name"],
                "sp": parsed["stored_path"], "tv": template["version"],
                "rd": json.dumps(recon), "rp": month}, cfg=cfg)
    ch_node, rp_node = f"CH_{ref}", f"RP_{month}"
    v = template["version"]
    nodes = [
        {"node_id": ch_node, "label": "Coverholder",
         "properties": json.dumps({"name": name, "coverholder_reference": ref}),
         "source_document_id": doc_id, "extraction_path": "det", "confidence": 1.0,
         "template_version": v, "dictionary_version": "v1",
         "anchor": json.dumps({"cell": "Bordereau!A1"}), "tenant": "bordereaux"},
        {"node_id": rp_node, "label": "ReportingPeriod",
         "properties": json.dumps({"reporting_month": month}),
         "source_document_id": doc_id, "extraction_path": "det", "confidence": 1.0,
         "template_version": v, "dictionary_version": "v1",
         "anchor": json.dumps({"cell": "Bordereau!A3"}), "tenant": "bordereaux"},
    ]
    return {"name": name, "month": month, "ch_node": ch_node, "rp_node": rp_node,
            "nodes": nodes, "superseded": superseded, "v": v, "ds": _short(doc_id)}


def _edge(src, tgt, rel, doc_id, v):
    return {"edge_id": f"E-{uuid.uuid4().hex[:10]}", "source_id": src, "target_id": tgt,
            "relationship_type": rel, "source_document_id": doc_id, "extraction_path": "det",
            "confidence": 1.0, "template_version": v, "dictionary_version": "v1",
            "anchor": None, "tenant": "bordereaux"}


def _line_node(node_id, label, props, doc_id, v, rownum, ncols):
    anchor = {"cell": f"Bordereau!A{rownum}:{_col_letter(ncols)}{rownum}"}
    return {"node_id": node_id, "label": label, "properties": json.dumps(props, default=str),
            "source_document_id": doc_id, "extraction_path": "reconciled", "confidence": 1.0,
            "template_version": v, "dictionary_version": "v1", "anchor": json.dumps(anchor),
            "tenant": "bordereaux"}


def land_premium(cfg, parsed, template, det_map, mapped_rows, recon, doc_id):
    F = cfg["full_schema"]
    lc = _land_common(cfg, parsed, template, recon, doc_id)
    node_rows, edge_rows, proj_rows = lc["nodes"], [], []
    ncols = len(parsed["headers"])
    for i, m in enumerate(mapped_rows):
        node = f"PBL_{lc['ds']}_{i+1:03d}"
        node_rows.append(_line_node(node, "PremiumBordereauLine", m, doc_id, lc["v"], HEADER_ROW + 1 + i, ncols))
        edge_rows.append(_edge(node, lc["ch_node"], "SUBMITTED_BY", doc_id, lc["v"]))
        edge_rows.append(_edge(node, lc["rp_node"], "FOR_PERIOD", doc_id, lc["v"]))
        proj_rows.append({
            "bordereau_line_id": node, "reporting_month": _month_date(lc["month"]),
            "coverholder_name": lc["name"], "policy_number": m.get("policy_number"),
            "insured_name": m.get("insured_name"),
            "line_of_business_code": LOB_CODE.get(m.get("line_of_business")),
            "inception_date": m.get("inception_date"), "expiry_date": m.get("expiry_date"),
            "currency_code": m.get("currency"), "gross_premium": m.get("gross_premium"),
            "commission_amount": m.get("commission_amount"),
            "risk_country_code": m.get("risk_country"), "risk_postcode": m.get("risk_postcode"),
            "source_system_code": "COVERHOLDER_BDX", "source_document_id": doc_id})
    proj_cols = [("bordereau_line_id", "str"), ("reporting_month", "date"),
                 ("coverholder_name", "str"), ("policy_number", "str"), ("insured_name", "str"),
                 ("line_of_business_code", "str"), ("inception_date", "date"),
                 ("expiry_date", "date"), ("currency_code", "str"), ("gross_premium", "dec"),
                 ("commission_amount", "dec"), ("risk_country_code", "str"),
                 ("risk_postcode", "str"), ("source_system_code", "str"), ("source_document_id", "str")]
    insert_rows(f"{F}.graph_nodes", NODE_COLS, node_rows, cfg)
    insert_rows(f"{F}.graph_edges", EDGE_COLS, edge_rows, cfg)
    insert_rows(f"{F}.premium_bordereau_line", proj_cols, proj_rows, cfg)
    audit(cfg, "ingested", "source_document", doc_id,
          detail=f"{lc['name']} premium {lc['month']}: {len(proj_rows)} lines landed")
    return {"state": "known_unchanged", "lines": len(proj_rows), "superseded": lc["superseded"]}


def land_claims(cfg, parsed, template, det_map, mapped_rows, recon, doc_id):
    F = cfg["full_schema"]
    lc = _land_common(cfg, parsed, template, recon, doc_id)
    node_rows, edge_rows, proj_rows = lc["nodes"], [], []
    ncols = len(parsed["headers"])
    for i, m in enumerate(mapped_rows):
        node = f"CBL_{lc['ds']}_{i+1:03d}"
        node_rows.append(_line_node(node, "ClaimBordereauLine", m, doc_id, lc["v"], HEADER_ROW + 1 + i, ncols))
        edge_rows.append(_edge(node, lc["ch_node"], "CLAIM_SUBMITTED_BY", doc_id, lc["v"]))
        edge_rows.append(_edge(node, lc["rp_node"], "FOR_PERIOD", doc_id, lc["v"]))
        paid = m.get("paid_amount") or 0
        outstanding = m.get("outstanding_amount") or 0
        try:
            incurred = round(float(paid) + float(outstanding), 2)
        except (TypeError, ValueError):
            incurred = None
        proj_rows.append({
            "claim_line_id": node, "reporting_month": _month_date(lc["month"]),
            "coverholder_name": lc["name"], "claim_reference": m.get("claim_reference"),
            "policy_number": m.get("policy_number"), "loss_date": m.get("loss_date"),
            "cause_of_loss_code": CAUSE_CODE.get(m.get("cause_of_loss")),
            "paid_amount": paid, "outstanding_amount": outstanding, "incurred_amount": incurred,
            "currency_code": m.get("currency"), "source_system_code": "COVERHOLDER_BDX",
            "source_document_id": doc_id})
    proj_cols = [("claim_line_id", "str"), ("reporting_month", "date"), ("coverholder_name", "str"),
                 ("claim_reference", "str"), ("policy_number", "str"), ("loss_date", "date"),
                 ("cause_of_loss_code", "str"), ("paid_amount", "dec"), ("outstanding_amount", "dec"),
                 ("incurred_amount", "dec"), ("currency_code", "str"), ("source_system_code", "str"),
                 ("source_document_id", "str")]
    insert_rows(f"{F}.graph_nodes", NODE_COLS, node_rows, cfg)
    insert_rows(f"{F}.graph_edges", EDGE_COLS, edge_rows, cfg)
    insert_rows(f"{F}.claim_bordereau_line", proj_cols, proj_rows, cfg)
    audit(cfg, "ingested", "source_document", doc_id,
          detail=f"{lc['name']} claims {lc['month']}: {len(proj_rows)} lines landed")
    return {"state": "known_unchanged", "lines": len(proj_rows), "superseded": lc["superseded"]}


def land_generic(cfg, parsed, template, det_map, mapped_rows, recon, doc_id, fam_cfg, tenant):
    """Config-driven lander for ANY tabular tenant. No family-specific code: the node label,
    projection table and per-column mapping all come from the tenant config. This is what makes
    onboarding a new document family (e.g. the health renewal slip) a config change, not code."""
    F = cfg["full_schema"]
    name, ref, period, family = parsed["name"], parsed["ref"], parsed["month"], parsed["family"]
    tn = tenant["tenant"]
    superseded = _supersede(cfg, name, family, period)
    sql(f"""INSERT INTO {F}.source_document VALUES
            (:doc,:tn,:c,:f,:fn,:sp,'active','known_unchanged',:tv,'v1',:rd,:rp,
             current_timestamp(),NULL,NULL)""",
        params={"doc": doc_id, "tn": tn, "c": name, "f": family, "fn": parsed["file_name"],
                "sp": parsed["stored_path"], "tv": template["version"],
                "rd": json.dumps(recon), "rp": period}, cfg=cfg)

    ds, v = _short(doc_id), template["version"]
    party, per = f"PARTY_{ref}", f"RP_{period}"
    node_rows = [
        {"node_id": party, "label": "Party",
         "properties": json.dumps({"name": name, "reference": ref}),
         "source_document_id": doc_id, "extraction_path": "det", "confidence": 1.0,
         "template_version": v, "dictionary_version": "v1",
         "anchor": json.dumps({"cell": "Sheet1!A1"}), "tenant": tn},
        {"node_id": per, "label": "ReportingPeriod",
         "properties": json.dumps({"reporting_period": period}),
         "source_document_id": doc_id, "extraction_path": "det", "confidence": 1.0,
         "template_version": v, "dictionary_version": "v1",
         "anchor": json.dumps({"cell": "Sheet1!A3"}), "tenant": tn},
    ]
    label = fam_cfg.get("node_label", "Record")
    proj = tenant["projections"][fam_cfg["projection"]]
    ncols = len(parsed["headers"])
    edge_rows, proj_rows = [], []
    specials = {"__doc_id": doc_id, "__period": period, "__party": name}
    for i, m in enumerate(mapped_rows):
        node = f"GEN_{ds}_{i+1:03d}"
        n = _line_node(node, label, m, doc_id, v, HEADER_ROW + 1 + i, ncols); n["tenant"] = tn
        node_rows.append(n)
        for tgt, rel in ((party, "REPORTED_BY"), (per, "FOR_PERIOD")):
            e = _edge(node, tgt, rel, doc_id, v); e["tenant"] = tn
            edge_rows.append(e)
        row = {}
        for c in proj["columns"]:
            frm = c["from"]
            row[c["col"]] = node if frm == "__node_id" else specials.get(frm, m.get(frm))
        proj_rows.append(row)

    proj_cols = [(c["col"], c["type"]) for c in proj["columns"]]
    insert_rows(f"{F}.graph_nodes", NODE_COLS, node_rows, cfg)
    insert_rows(f"{F}.graph_edges", EDGE_COLS, edge_rows, cfg)
    insert_rows(f"{F}.{proj['table']}", proj_cols, proj_rows, cfg)
    audit(cfg, "ingested", "source_document", doc_id,
          detail=f"{name} {family} {period}: {len(proj_rows)} rows landed (generic lander)")
    return {"state": "known_unchanged", "lines": len(proj_rows), "superseded": superseded}


LANDERS = {"premium": land_premium, "claims": land_claims}  # 'generic' handled explicitly


def quarantine(cfg, parsed, doc_id, inbound_state, template, recon, reason=""):
    F = cfg["full_schema"]
    tv = template["version"] if template else None
    detail = recon if recon else {"reason": reason}
    sql(f"""INSERT INTO {F}.source_document VALUES
            (:doc,'bordereaux',:c,:f,:fn,:sp,'quarantined',:st,:tv,'v1',:rd,:rp,
             current_timestamp(),NULL,NULL)""",
        params={"doc": doc_id, "c": parsed["name"], "f": parsed["family"],
                "fn": parsed["file_name"], "sp": parsed["stored_path"], "st": inbound_state,
                "tv": tv, "rd": json.dumps(detail), "rp": parsed["month"]}, cfg=cfg)
    audit(cfg, "quarantined", "source_document", doc_id,
          detail=f"{parsed['name']} {parsed['family']} {parsed['month']}: {inbound_state} — {reason}")
    return {"state": inbound_state, "lines": 0}


# ------------------------------------------------------------------- orchestration
def process_file(cfg, path, tenant=None):
    """Ingest one document. `tenant` is the tenant config dict (loaded if omitted).
    Branches on document family for the correct logical fields + landing projection."""
    tenant = tenant or load_tenant(cfg, "bordereaux")
    doc_id = f"DOC-{uuid.uuid4().hex[:12]}"
    fname = os.path.basename(path)
    base = {"file_name": fname, "stored_path": path, "name": fname, "ref": "",
            "month": "", "family": "premium_bordereau"}
    try:
        parsed = parse_file(cfg, path)
    except Unusable as e:
        return quarantine(cfg, {**base}, doc_id, "unusable", None, None, reason=str(e))

    parsed["file_name"] = fname
    parsed["stored_path"] = path
    # Resolve the family + its config from the filename token, config-driven (no per-family code).
    token = parsed.get("family_token")
    fam_name, fam_cfg = None, None
    for fn, fv in tenant.get("families", {}).items():
        if fv.get("file_token") == token or fn == parsed["family"]:
            fam_name, fam_cfg = fn, fv
            break
    if fam_cfg is None:
        fam_name, fam_cfg = parsed["family"], {}
    parsed["family"] = fam_name
    logical_fields = tenant["logical_fields"].get(fam_cfg.get("fields", ""), [])
    template = lookup_template(cfg, parsed["name"], fam_name)
    if template is None:
        recon = {"reason": "no template for counterparty — schema recognition required",
                 "headers": parsed["headers"]}
        return quarantine(cfg, parsed, doc_id, "unknown", None, recon, reason="unknown counterparty")

    det_map, mapped_rows = det_extract(template["field_map"], parsed["headers"], parsed["rows"])
    ai_map = ai_extract(cfg, parsed["headers"], parsed["rows"], logical_fields)
    recon = reconcile(cfg, det_map, ai_map, logical_fields)

    all_present = all(v["present"] for v in det_map.values())
    if all_present and recon["summary"]["differs"] == 0:
        lname = fam_cfg.get("lander", "generic")
        if lname == "generic":
            return land_generic(cfg, parsed, template, det_map, mapped_rows, recon, doc_id, fam_cfg, tenant)
        return LANDERS[lname](cfg, parsed, template, det_map, mapped_rows, recon, doc_id)
    return quarantine(cfg, parsed, doc_id, "known_changed", template, recon,
                      reason=f"{recon['summary']['differs']} field(s) differ / "
                             f"{recon['summary']['expected'] - recon['summary']['present']} missing")


def process_inbox(cfg=None, tenant_name="bordereaux", subdir="inbox"):
    cfg = cfg or load_config()
    tenant = load_tenant(cfg, tenant_name)
    base = f'{cfg["volume_path"]}/{subdir}'
    try:
        files = [f.path for f in client(cfg).files.list_directory_contents(base)]
    except Exception:
        print(f"== no inbox at {base}")
        return {}
    print(f"== processing {len(files)} {tenant_name} files in {subdir}")
    summary = {}
    for path in sorted(files):
        res = process_file(cfg, path, tenant)
        state = res["state"]
        summary[state] = summary.get(state, 0) + 1
        print(f"  {os.path.basename(path):<40} -> {state} ({res.get('lines',0)} lines)")
    print("== inbox summary:", summary)
    return summary


if __name__ == "__main__":
    process_inbox()
